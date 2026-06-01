import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path


def main() -> None:
    rows = [
        [
            1,
            "Baseline scan of Tableau files (.twb/.twbx/.tds/.tdsx)",
            "",
            "Completed",
            "Source scan works for local folders and lists convertible temp-table queries",
            "",
        ],
        [
            2,
            "Extract Custom SQL and Initial/One-time SQL from Tableau XML",
            "",
            "Completed",
            'Parser reads <relation type="text"> and initial-sql/one-time-sql patterns',
            "",
        ],
        [
            3,
            "Temp-table detection + dependency ordering (multi-statement SQL)",
            "",
            "Completed",
            "Dependency graph handles multi-level temp table chains (with sqlglot + fallback)",
            "",
        ],
        [
            4,
            "Temp -> permanent table conversion logic",
            "",
            "Completed",
            "Conversion produces permanent-table version SQL and analysis",
            "",
        ],
        [
            5,
            "Write-back converted SQL into workbook and download",
            "",
            "Completed",
            "Upload-based write-back works; backend also supports write-back from scan results",
            "",
        ],
        [
            6,
            "Main issue: no visibility of converted permanent table in Tableau datasource",
            "",
            "In Progress",
            "After write-back, Tableau sometimes does not show expected fields/metadata; likely contract/placement mismatch",
            "",
        ],
        [
            7,
            "Main issue: semicolon / multi-statement handling in Custom SQL and Initial SQL",
            "",
            "In Progress",
            "Tableau often expects single-statement Custom SQL; current converted SQL may include multi statements",
            "",
        ],
        [
            8,
            "Main issue: Tableau conversion contract (what runs at setup vs runtime)",
            "",
            "In Progress",
            "Need a fixed approach: setup SQL location + runtime query shape to reduce breakage",
            "",
        ],
        [
            9,
            "Main issue: multi-level temp tables in specific workbooks (edge cases)",
            "",
            "In Progress",
            "Needs targeted test workbooks and stricter statement classification for unusual patterns",
            "",
        ],
        [
            10,
            "Main issue: high-end query optimization",
            "",
            "In Progress",
            "Permanent tables can add cost/locking; need performance-friendly strategy (naming/cleanup/indexing)",
            "",
        ],
        [
            11,
            "Improve query formatting + stable matching during write-back",
            "",
            "In Progress",
            "Better normalization/formatting reduces mismatches and improves readability",
            "",
        ],
        [
            12,
            "Parallel work: review other git repos for better .twbx patching patterns",
            "",
            "In Progress",
            "Comparing approaches for patch accuracy across Tableau XML variations",
            "",
        ],
        [
            13,
            "Parallel work: review other git repos for SQL formatting/refactor patterns",
            "",
            "In Progress",
            "Looking for safer handling of multi-statement SQL and Tableau constraints",
            "",
        ],
        [
            14,
            "Next deliverable: compatibility visibility check after conversion",
            "",
            "Pending",
            "Add report for missing/renamed output columns and likely broken fields before download",
            "",
        ],
        [
            15,
            "Next deliverable: updated write-back approach aligned to Tableau rules",
            "",
            "Pending",
            "Ensure runtime query is Tableau-compatible; move setup SQL appropriately",
            "",
        ],
    ]

    headers = ["S.No", "Topics", "Owner", "Status", "Comments", "Target Date"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Status"

    ws.append(headers)
    for row in rows:
        ws.append(row)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6F63")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wrap_cols = {2, 5}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in wrap_cols:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top")

    widths = {
        1: 6,
        2: 70,
        3: 18,
        4: 14,
        5: 90,
        6: 14,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    out_path = Path(r"A:\analytics_governance\tableau_poc_status.xlsx")
    wb.save(out_path)


if __name__ == "__main__":
    main()

